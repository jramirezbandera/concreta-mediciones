"""Extrae el modelo de estimación de la hoja «-RCDs =+ ECO-eficiencia» (ECO-ARQ, Universidad
de Sevilla; 20130712RCDs.xlsm) a `src/core/rcd/tablas/ecoeficiencia-2013.json`.

La hoja no está en el repo (documentación de apoyo del COA de Málaga, con derechos de los
autores y de la Consejería): se pasa su ruta. Lee las FÓRMULAS, no los valores calculados:

  · nivel detallado (hoja PASO2.D): ratio base por código (columna D, m³ aparentes / m²
    construido), coeficientes por opción (tipología, plantas, cimentación, estructura vertical
    y horizontal, cubierta, fachada, pavimento) y sumandos por desviación de superficies
    respecto al edificio de referencia (columnas AZ–BF: «=P4x*f*K9»);
  · tierras (filas de limpieza, excavación y vaciado);
  · nivel básico (hoja PASO2.B): reparto por código de un ratio global por tipo de obra;
  · densidad aparente por código (hoja PASO3: «=E10*0.07»).

Uso: python -I scripts/rcd/extraer_ecoeficiencia.py <20130712RCDs.xlsm> <salida.json>
Cualquier celda que no encaje con lo esperado sale por pantalla y aborta.
"""
import json
import re
import sys

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

sys.stdout.reconfigure(encoding='utf-8')
src, out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src, data_only=False)
D2 = wb['PASO2.D']
B2 = wb['PASO2.B']
P3 = wb['PASO3']
problemas = []


def v(ws, ref):
    return ws[ref].value


def num(x, donde):
    if isinstance(x, (int, float)):
        return float(x)
    if isinstance(x, str):
        m = re.fullmatch(r'=?\s*(-?\d+(?:\.\d+)?)', x.strip())
        if m:
            return float(m.group(1))
    problemas.append(f'{donde}: no es un número: {x!r}')
    return None


def codigo(x):
    return re.sub(r'\s+', ' ', str(x)).strip()


def nombre(x):
    return re.sub(r'\s+', ' ', str(x)).replace('*', '').strip()


# ---- opciones (cabecera de la tabla de coeficientes, fila 52) ----------------------------
GRUPOS = [
    ('tipologia', 'G', 5, ['plurifamiliarExento', 'plurifamiliarEntreMedianeras', 'unifamiliarExento',
                           'unifamiliarEntreMedianeras', 'unifamiliarPareado']),
    ('plantas', 'M', 5, ['1', '2-3', '4-5', '6-8', '>8']),
    ('cimentacion', 'S', 5, ['zapatas', 'losa', 'pilotesInSitu', 'pilotesPrefabricados', 'murosSotano']),
    ('estructuraVertical', 'Y', 4, ['hormigonArmado', 'acero', 'murosCarga', 'otros']),
    ('estructuraHorizontal', 'AD', 4, ['unidireccionalBovedillaCeramica', 'unidireccionalBovedillaMortero',
                                       'reticular', 'otros']),
    ('cubierta', 'AI', 4, ['planaAndaluza', 'planaInvertida', 'inclinadaTeja', 'otros']),
    ('fachada', 'AN', 6, ['ceramicaCaraVista', 'ceramicaRevestida', 'bloqueCaraVista', 'bloqueRevestido',
                          'ventiladaAplacado', 'otros']),
    ('pavimento', 'AU', 4, ['piedraNatural', 'piedraArtificial', 'baldosaCeramica', 'otros']),
]
etiquetas = {}
for g, col, n, ids in GRUPOS:
    c0 = column_index_from_string(col)
    etiquetas[g] = []
    for i in range(n):
        lab = v(D2, f'{get_column_letter(c0 + i)}52')
        etiquetas[g].append({'id': ids[i], 'hoja': str(lab).strip()})

# ---- edificio de referencia (O44..O49: «=37/2882»…) --------------------------------------
REF = {}
for fila, clave in [(44, 'arquetas'), (45, 'forjados'), (46, 'cubiertas'), (47, 'hojaExterior'),
                    (48, 'particiones'), (49, 'pavimentos')]:
    f = str(v(D2, f'O{fila}'))
    m = re.fullmatch(r'=(\d+)/(\d+)', f)
    if not m:
        problemas.append(f'O{fila}: referencia inesperada {f!r}')
        continue
    REF[clave] = int(m.group(1))
    REF['superficieConstruida'] = int(m.group(2))
SUMANDO_COL = {'AZ': ('arquetas', 44), 'BA': ('forjados', 45), 'BB': ('cubiertas', 46),
               'BC': ('hojaExterior', 47), 'BD': ('particiones', 48), 'BE': ('pavimentos', 49)}

# ---- densidad aparente por fila de PASO2.D, desde PASO3 («=PASO2.B!E45+PASO2.D!F54» y su
# vecina «=E10*0.07») --------------------------------------------------------------------------
dens_por_fila = {}
for row in P3.iter_rows():
    for c in row:
        f = c.value
        if isinstance(f, str):
            m = re.fullmatch(r'=PASO2\.B!E(\d+)\+PASO2\.D!F(\d+)', f.replace(' ', ''))
            if m:
                vec = P3.cell(row=c.row, column=c.column + 1).value
                if re.fullmatch(r'=([A-Z]+\d+\+)+[A-Z]+\d+', str(vec).replace(' ', '')):
                    continue  # fila de grupo (envases, restos, tierras): es una suma
                mm = re.fullmatch(rf'={get_column_letter(c.column)}{c.row}\*(\d+(?:\.\d+)?)', str(vec))
                if not mm:
                    problemas.append(f'PASO3 {get_column_letter(c.column + 1)}{c.row}: densidad {vec!r}')
                    continue
                dens_por_fila[int(m.group(2))] = (float(mm.group(1)), int(m.group(1)))

# ---- filas del nivel detallado -------------------------------------------------------------
GRUPO_FILA = {}
filas = []
grupo_actual = None
for r in range(53, 86):
    cod = v(D2, f'B{r}')
    if cod is None:
        continue
    cod = codigo(cod)
    nom = nombre(v(D2, f'C{r}'))
    if cod in ('15', '17', '17 05'):
        grupo_actual = {'15': 'envases', '17': 'restos', '17 05': 'tierras'}[cod]
        continue
    if grupo_actual == 'tierras' and nom.lower().startswith('tierras'):
        continue  # limpieza, excavación y vaciado: fórmulas propias (abajo)
    base = num(v(D2, f'D{r}'), f'D{r}')
    fF = str(v(D2, f'F{r}')).replace(' ', '')
    esperado = f'=E{r}*L{r}*R{r}*X{r}*AC{r}*AH{r}*AM{r}*AT{r}*AY{r}+BF{r}'
    if fF != esperado:
        problemas.append(f'F{r}: fórmula inesperada {fF!r}')
    if str(v(D2, f'E{r}')).replace(' ', '') != f'=K9*D{r}':
        problemas.append(f'E{r}: fórmula inesperada {v(D2, f"E{r}")!r}')
    coef = {}
    for g, col, n, _ in GRUPOS:
        c0 = column_index_from_string(col)
        vals = []
        for i in range(n):
            x = D2.cell(row=r, column=c0 + i).value
            vals.append(0.0 if x is None else num(x, f'{get_column_letter(c0 + i)}{r}'))
        coef[g] = vals
    # sumandos: los que BF suma
    bf = str(v(D2, f'BF{r}') or '').replace(' ', '')
    usados = re.findall(r'([A-Z]{2})' + str(r), bf) if bf else []
    sumandos = {}
    for colS in usados:
        f = str(v(D2, f'{colS}{r}')).replace(' ', '')
        clave, fila_p = SUMANDO_COL[colS]
        m = re.fullmatch(rf'=P{fila_p}\*(\d+(?:\.\d+)?)\*K9', f)
        if not m:
            problemas.append(f'{colS}{r}: sumando inesperado {f!r}')
            continue
        sumandos[clave] = float(m.group(1))
    if r not in dens_por_fila:
        problemas.append(f'fila {r} ({cod} {nom}): sin densidad en PASO3')
    dens, fila_b = dens_por_fila.get(r, (None, None))
    filas.append({'ler': cod, 'nombre': nom, 'grupo': grupo_actual, 'm3PorM2': base,
                  'densidadAparente': dens, 'coef': coef, 'sumandos': sumandos, '_filaB': fila_b})

# ---- tierras (filas 80–82) -----------------------------------------------------------------
def formula(ws, ref):
    return str(ws[ref].value).replace(' ', '')


t80, t81, t82 = formula(D2, 'E80'), formula(D2, 'E81'), formula(D2, 'E82')
m80 = re.fullmatch(r'=K13\*(\d+(?:\.\d+)?)', t80)
m81 = re.fullmatch(r'=K15\*(\d+(?:\.\d+)?)', t81)
m82 = re.fullmatch(r'=E17\*(\d+(?:\.\d+)?)\*K15', t82)
if not (m80 and m81 and m82):
    problemas.append(f'tierras: {t80!r} {t81!r} {t82!r}')
if formula(D2, 'F81') != '=E81*X81+BF81' or formula(D2, 'BF81') != '=AZ81':
    problemas.append(f'tierras F81/BF81: {formula(D2, "F81")!r} {formula(D2, "BF81")!r}')
m_arq = re.fullmatch(r'=P44\*(\d+(?:\.\d+)?)\*K9', formula(D2, 'AZ81'))
coef_exc = []
for i in range(5):
    x = D2.cell(row=81, column=column_index_from_string('S') + i).value
    coef_exc.append(1.0 if x is None and i == 0 else (0.0 if x is None else num(x, f'S81+{i}')))
tierras = {
    'limpiezaDesbroce': {'ler': '17 05 04', 'base': 'superficieSolar', 'm3PorM2': float(m80.group(1)),
                         'densidadAparente': dens_por_fila.get(80, (None,))[0]},
    'excavacion': {'ler': '17 05 04', 'base': 'superficieOcupacion', 'm3PorM2': float(m81.group(1)),
                   'coefCimentacion': coef_exc, 'arquetas': float(m_arq.group(1)) if m_arq else None,
                   'densidadAparente': dens_por_fila.get(81, (None,))[0]},
    'vaciadoSotano': {'ler': '17 05 04', 'base': 'superficieOcupacion', 'alturaPorPlanta': float(m82.group(1)),
                      'densidadAparente': dens_por_fila.get(82, (None,))[0]},
}

# ---- nivel básico (PASO2.B, columnas F–K) --------------------------------------------------
TIPOS_B = [('F', 'nuevaResidencial', 'E28'), ('G', 'nuevaNoResidencial', 'E28'),
           ('H', 'demolicionResidencial', 'E30'), ('I', 'demolicionNoResidencial', 'E30'),
           ('J', 'rehabilitacionResidencial', 'E32'), ('K', 'rehabilitacionNoResidencial', 'E32')]
coef_plantas_b = [num(v(B2, f'{c}45'), f'{c}45') for c in ('M', 'N', 'O')]
basico = []
for r in list(range(45, 50)) + list(range(52, 69)) + list(range(74, 77)):
    cod = codigo(v(B2, f'B{r}'))
    fila = {'ler': cod, 'nombre': nombre(v(B2, f'C{r}')), 'porTipo': {}}
    for col, tipo, base in TIPOS_B:
        f = v(B2, f'{col}{r}')
        f = '0' if f is None else str(f).replace(' ', '')
        if f in ('0', '=0'):
            fila['porTipo'][tipo] = 0.0
            continue
        m = re.fullmatch(rf'=(\d+(?:\.\d+)?)\*{base}\*(\d+(?:\.\d+)?)', f)
        if m:
            fila['porTipo'][tipo] = round(float(m.group(1)) * float(m.group(2)), 10)
            continue
        m = re.fullmatch(rf'=(\d+(?:\.\d+)?)\*{base}', f)
        if m:
            fila['porTipo'][tipo] = float(m.group(1))
            continue
        m = re.fullmatch(rf'=IF\(M34=1,(\d+(?:\.\d+)?)\*{base}(?:\*(\d+(?:\.\d+)?))?,(\d+(?:\.\d+)?)\*{base}(?:\*(\d+(?:\.\d+)?))?\)', f)
        if m:
            k1 = float(m.group(2) or 1); k2 = float(m.group(4) or 1)
            fila['porTipo'][tipo] = {'estructuraHormigon': round(float(m.group(1)) * k1, 10),
                                     'otraEstructura': round(float(m.group(3)) * k2, 10)}
            continue
        problemas.append(f'PASO2.B {col}{r}: fórmula inesperada {f!r}')
    fila['densidadAparente'] = next((d for fr, (d, fb) in dens_por_fila.items() if fb == r), None)
    if fila['densidadAparente'] is None:
        problemas.append(f'PASO2.B fila {r} ({cod}): sin densidad')
    fE = str(v(B2, f'E{r}')).replace(' ', '')
    if fE == f'=L{r}*P{r}':
        fila['coefPlantas'] = True
    elif fE == f'=L{r}':
        fila['coefPlantas'] = False
    else:
        problemas.append(f'PASO2.B E{r}: fórmula inesperada {fE!r}')
    basico.append(fila)
# tierras del nivel básico (filas 71–73): limpieza y excavación sobre el solar, vaciado tecleado
b71, b72, b73 = formula(B2, 'E71'), formula(B2, 'E72'), formula(B2, 'E73')
mb71 = re.fullmatch(r'=IF\(E28=0,0,E34\*(\d+(?:\.\d+)?)\)', b71)
mb72 = re.fullmatch(r'=IF\(E28=0,0,E34\*(\d+(?:\.\d+)?)\)', b72)
if not (mb71 and mb72 and b73 == '=I32'):
    problemas.append(f'tierras básicas: {b71!r} {b72!r} {b73!r}')
tierras_basico = {
    'nota': 'Solo con superficie construida (obra nueva); sobre la superficie del solar',
    'limpiezaDesbroce': {'ler': '17 05 04', 'm3PorM2Solar': float(mb71.group(1)) if mb71 else None,
                         'densidadAparente': dens_por_fila.get(80, (None,))[0]},
    'excavacion': {'ler': '17 05 04', 'm3PorM2Solar': float(mb72.group(1)) if mb72 else None,
                   'densidadAparente': dens_por_fila.get(81, (None,))[0]},
    'vaciadoSotano': {'ler': '17 05 04', 'base': 'volumen de vaciado tecleado (m³)',
                      'densidadAparente': dens_por_fila.get(82, (None,))[0]},
}
for fila in filas:
    fila.pop('_filaB', None)

if problemas:
    print('\n'.join(problemas))
    sys.exit(1)

doc = {
    'id': 'ecoeficiencia-2013',
    'fuente': '«-RCDs = +ECO-eficiencia. Herramienta de ayuda a la reducción de residuos en el diseño '
              'y construcción de viviendas en Andalucía»',
    'autor': 'Equipo de investigación ECO-ARQ, Departamento de Construcciones Arquitectónicas I, '
             'Escuela Técnica Superior de Arquitectura, Universidad de Sevilla',
    'edicion': 'Manual de octubre de 2012; hoja de cálculo 20130712RCDs.xlsm (12-07-2013)',
    'subvencion': 'Consejería de Fomento y Vivienda de la Junta de Andalucía (Orden de 3-03-2009)',
    'metodo': 'Llatas, C. (2011). A model for quantifying construction waste in projects according to '
              'the European waste list. Waste Management, 31(6), 1261-1276.',
    'derechos': 'El manual (ap. 1.5) no permite la difusión, comercialización o reproducción sin '
                'autorización escrita de la Consejería y de los autores. Uso provisional, pendiente de '
                'permiso (plan D8).',
    'ambito': 'Vivienda de nueva planta en Andalucía (obras de VPO 2009-2012); nivel básico estimativo '
              'para otros usos, rehabilitación y demolición.',
    'unidad': 'm³ aparentes por m² (de superficie construida, de solar, de ocupación, demolida o '
              'rehabilitada según la fila); densidad aparente en t/m³',
    'referencia': REF,
    'opciones': etiquetas,
    'detallado': filas,
    'tierras': tierras,
    'basico': {'coefPlantas': {'opciones': ['1', '2-4', '>4'], 'valores': coef_plantas_b}, 'filas': basico,
               'tierras': tierras_basico},
}
json.dump(doc, open(out, 'w', encoding='utf-8', newline='\n'), ensure_ascii=False, indent=1)
print('ok:', len(filas), 'filas detalladas,', len(basico), 'filas básicas')
